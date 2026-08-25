"""Export der Zeiterfassung als formatierte Excel-Arbeitsmappe.

Drei Blaetter:

``Uebersicht``  Kopfzahlen des Zeitraums und die Wochensalden.
``Tage``        Eine Zeile je Tag mit Kommen, Gehen, Pause, Ist, Soll, Saldo.
``Buchungen``   Jede einzelne Sitzung und jede eingeordnete Luecke -- der
                Nachweis hinter den Tageszahlen.
"""

from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .rules import Tageswerte, als_dezimal, wochenbeginn
from .storage import ABWESEND, ARBEIT, PAUSE, Datenbank
from .tracker import Zeiterfassung

WOCHENTAGE = ("Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag")
ART_TEXT = {ARBEIT: "Arbeit", PAUSE: "Pause", ABWESEND: "Abwesend"}

FARBE_KOPF = "1F3864"
FARBE_ZEILE = "F2F5FA"
FARBE_WOCHENENDE = "EFEFEF"
FARBE_PLUS = "1E7B34"
FARBE_MINUS = "B00020"

_KOPF_SCHRIFT = Font(color="FFFFFF", bold=True, size=11)
_KOPF_FUELLUNG = PatternFill("solid", fgColor=FARBE_KOPF)
_RAHMEN = Border(bottom=Side(style="thin", color="D0D5DD"))


def _kopfzeile(blatt: Worksheet, zeile: int, spalten: list[str]) -> None:
    for spalte, text in enumerate(spalten, start=1):
        zelle = blatt.cell(row=zeile, column=spalte, value=text)
        zelle.font = _KOPF_SCHRIFT
        zelle.fill = _KOPF_FUELLUNG
        zelle.alignment = Alignment(horizontal="center", vertical="center")
    blatt.freeze_panes = blatt.cell(row=zeile + 1, column=1)
    blatt.row_dimensions[zeile].height = 22


def _breiten(blatt: Worksheet, breiten: list[int]) -> None:
    for nummer, breite in enumerate(breiten, start=1):
        blatt.column_dimensions[get_column_letter(nummer)].width = breite


def _saldofarbe(zelle) -> None:
    wert = zelle.value
    if isinstance(wert, (int, float)):
        zelle.font = Font(color=FARBE_PLUS if wert >= 0 else FARBE_MINUS, bold=True)


def _uhrzeit(zeitpunkt: datetime | None) -> str:
    return zeitpunkt.strftime("%H:%M") if zeitpunkt else ""


def exportieren(
    zeiterfassung: Zeiterfassung,
    von: date,
    bis: date,
    ziel: Path | str,
    jetzt: datetime | None = None,
) -> Path:
    """Schreibt den Zeitraum als .xlsx und liefert den Pfad zurueck."""
    ziel = Path(ziel)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    tage = zeiterfassung.zeitraum(von, bis, jetzt=jetzt)

    mappe = Workbook()
    _blatt_uebersicht(mappe.active, tage, von, bis)
    _blatt_tage(mappe.create_sheet("Tage"), tage)
    _blatt_buchungen(mappe.create_sheet("Buchungen"), zeiterfassung.db, von, bis)
    mappe.save(ziel)
    return ziel


def _blatt_uebersicht(blatt: Worksheet, tage: list[Tageswerte], von: date, bis: date) -> None:
    blatt.title = "Uebersicht"
    _breiten(blatt, [28, 16, 16, 16, 16])

    blatt["A1"] = "Zeiterfassung"
    blatt["A1"].font = Font(size=18, bold=True, color=FARBE_KOPF)
    blatt["A2"] = f"Zeitraum {von:%d.%m.%Y} bis {bis:%d.%m.%Y}"
    blatt["A2"].font = Font(size=11, color="475467")

    ist = sum((t.arbeitszeit for t in tage), timedelta(0))
    soll = sum((t.soll for t in tage), timedelta(0))
    pause = sum((t.erfasste_pause + t.pausenabzug for t in tage), timedelta(0))
    arbeitstage = sum(1 for t in tage if t.arbeitszeit > timedelta(0))

    zeile = 4
    for beschriftung, wert in (
        ("Iststunden", als_dezimal(ist)),
        ("Sollstunden", als_dezimal(soll)),
        ("Saldo", als_dezimal(ist - soll)),
        ("Pausen", als_dezimal(pause)),
        ("Tage mit Erfassung", arbeitstage),
    ):
        blatt.cell(row=zeile, column=1, value=beschriftung).font = Font(bold=True)
        zelle = blatt.cell(row=zeile, column=2, value=wert)
        zelle.number_format = "0.00" if beschriftung != "Tage mit Erfassung" else "0"
        if beschriftung == "Saldo":
            _saldofarbe(zelle)
        zeile += 1

    zeile += 1
    _kopfzeile(blatt, zeile, ["Kalenderwoche", "Ist", "Soll", "Saldo"])
    blatt.freeze_panes = None
    erste_datenzeile = zeile + 1

    wochen: dict[date, list[Tageswerte]] = {}
    for tag in tage:
        wochen.setdefault(wochenbeginn(tag.tag), []).append(tag)

    for montag, wochentage in sorted(wochen.items()):
        zeile += 1
        w_ist = sum((t.arbeitszeit for t in wochentage), timedelta(0))
        w_soll = sum((t.soll for t in wochentage), timedelta(0))
        blatt.cell(row=zeile, column=1, value=f"KW {montag.isocalendar().week} (ab {montag:%d.%m.})")
        for spalte, wert in ((2, w_ist), (3, w_soll), (4, w_ist - w_soll)):
            zelle = blatt.cell(row=zeile, column=spalte, value=als_dezimal(wert))
            zelle.number_format = "0.00"
        _saldofarbe(blatt.cell(row=zeile, column=4))

    if zeile >= erste_datenzeile:
        diagramm = BarChart()
        diagramm.title = "Ist- und Sollstunden je Woche"
        diagramm.y_axis.title = "Stunden"
        diagramm.height, diagramm.width = 7, 16
        daten = Reference(blatt, min_col=2, max_col=3, min_row=erste_datenzeile - 1, max_row=zeile)
        beschriftung = Reference(blatt, min_col=1, min_row=erste_datenzeile, max_row=zeile)
        diagramm.add_data(daten, titles_from_data=True)
        diagramm.set_categories(beschriftung)
        blatt.add_chart(diagramm, f"F{erste_datenzeile - 1}")


def _blatt_tage(blatt: Worksheet, tage: list[Tageswerte]) -> None:
    _breiten(blatt, [12, 12, 10, 10, 10, 10, 10, 10, 12, 30])
    _kopfzeile(
        blatt,
        1,
        ["Datum", "Wochentag", "Kommen", "Gehen", "Anwesend", "Pause", "Ist", "Soll", "Saldo", "Hinweis"],
    )

    for nummer, tag in enumerate(tage, start=2):
        wochenende = tag.tag.weekday() >= 5
        hinweise = []
        if tag.laeuft:
            hinweise.append("laeuft noch")
        if tag.pausenabzug > timedelta(0):
            hinweise.append("Pflichtpause abgezogen")

        werte = [
            tag.tag,
            WOCHENTAGE[tag.tag.weekday()],
            _uhrzeit(tag.erste_anmeldung),
            _uhrzeit(tag.letzter_kontakt),
            als_dezimal(tag.anwesenheit),
            als_dezimal(tag.erfasste_pause + tag.pausenabzug),
            als_dezimal(tag.arbeitszeit),
            als_dezimal(tag.soll),
            als_dezimal(tag.arbeitszeit - tag.soll),
            ", ".join(hinweise),
        ]
        for spalte, wert in enumerate(werte, start=1):
            zelle = blatt.cell(row=nummer, column=spalte, value=wert)
            zelle.border = _RAHMEN
            if spalte == 1:
                zelle.number_format = "DD.MM.YYYY"
            elif spalte in (5, 6, 7, 8, 9):
                zelle.number_format = "0.00"
                zelle.alignment = Alignment(horizontal="right")
            elif spalte in (3, 4):
                zelle.alignment = Alignment(horizontal="center")
            if wochenende:
                zelle.fill = PatternFill("solid", fgColor=FARBE_WOCHENENDE)
            elif nummer % 2 == 0:
                zelle.fill = PatternFill("solid", fgColor=FARBE_ZEILE)
        _saldofarbe(blatt.cell(row=nummer, column=9))

    letzte = len(tage) + 1
    summe = letzte + 1
    blatt.cell(row=summe, column=2, value="Summe").font = Font(bold=True)
    for spalte in (5, 6, 7, 8, 9):
        buchstabe = get_column_letter(spalte)
        zelle = blatt.cell(row=summe, column=spalte, value=f"=SUM({buchstabe}2:{buchstabe}{letzte})")
        zelle.number_format = "0.00"
        zelle.font = Font(bold=True)
    if tage:
        blatt.auto_filter.ref = f"A1:J{letzte}"


def _blatt_buchungen(blatt: Worksheet, datenbank: Datenbank, von: date, bis: date) -> None:
    _breiten(blatt, [12, 10, 10, 12, 14, 40])
    _kopfzeile(blatt, 1, ["Datum", "Von", "Bis", "Dauer", "Art", "Bemerkung"])

    fenster_von = datetime.combine(von, datetime.min.time())
    fenster_bis = datetime.combine(bis, datetime.min.time()) + timedelta(days=1)

    eintraege: list[tuple[datetime, datetime, str, str]] = []
    for sitzung in datenbank.sitzungen(fenster_von, fenster_bis):
        bemerkung = "Rechner eingeschaltet"
        if sitzung.laeuft:
            bemerkung += " -- laeuft noch"
        elif sitzung.ende_geschaetzt:
            bemerkung += " -- Ende aus letztem Herzschlag"
        eintraege.append((sitzung.beginn, sitzung.ende, "Rechnerlaufzeit", bemerkung))
    for luecke in datenbank.luecken(fenster_von, fenster_bis):
        eintraege.append((luecke.beginn, luecke.ende, ART_TEXT.get(luecke.art, luecke.art), luecke.notiz))

    for nummer, (beginn, ende, art, bemerkung) in enumerate(sorted(eintraege), start=2):
        werte = [
            beginn.date(),
            beginn.strftime("%H:%M"),
            ende.strftime("%H:%M"),
            als_dezimal(ende - beginn),
            art,
            bemerkung,
        ]
        for spalte, wert in enumerate(werte, start=1):
            zelle = blatt.cell(row=nummer, column=spalte, value=wert)
            zelle.border = _RAHMEN
            if spalte == 1:
                zelle.number_format = "DD.MM.YYYY"
            elif spalte == 4:
                zelle.number_format = "0.00"
            if nummer % 2 == 0:
                zelle.fill = PatternFill("solid", fgColor=FARBE_ZEILE)
    if eintraege:
        blatt.auto_filter.ref = f"A1:F{len(eintraege) + 1}"
