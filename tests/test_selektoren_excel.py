"""SAP-Feld-IDs als Arbeitsmappe aus- und wieder einlesen.

Wozu das gebraucht wird
-----------------------
Die Feld-IDs sind muehsam erarbeitet: aufgezeichnet, zugeordnet, am
Zielsystem kontrolliert.  Sie liegen in einer JSON-Datei neben den
Einstellungen -- gut fuer das Programm, unhandlich fuer alles andere.
Als Arbeitsmappe lassen sie sich sichern, weitergeben, auf einen zweiten
Rechner uebertragen und in Ruhe durchsehen.

Der heikle Punkt ist der Haken "Geprueft".  Er bedeutet: jemand hat
diese ID am echten System kontrolliert, die Anwendung darf damit
schreiben.  Ob das fuer *dieses* System gilt, weiss eine Mappe nicht --
sie kann aus einer anderen Anlage stammen.  Waere der Haken beliebig
importierbar, waere der Import ein Weg, die Freigabe zu umgehen.
Deshalb kommt er nur dort mit, wo die ID unveraendert bleibt.
"""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

_TEMP_HOME = tempfile.mkdtemp(prefix="sap_excel_")
os.environ["SAP_ANGEBOT_HOME"] = _TEMP_HOME

from app.sap.selectors import SelectorRegistry  # noqa: E402
from app.services.selektoren_excel import (  # noqa: E402
    SPALTEN,
    exportiere_selektoren,
    lies_selektoren,
    uebernimm_selektoren,
)


class ExportTest(unittest.TestCase):
    """Was in der Mappe steht."""

    def setUp(self):
        self.ordner = Path(tempfile.mkdtemp(prefix="export_"))
        self.registry = SelectorRegistry()

    def _mappe(self):
        from openpyxl import load_workbook

        pfad = exportiere_selektoren(self.registry, self.ordner / "ids.xlsx")
        return load_workbook(pfad).active

    def test_jedes_feld_steht_darin(self):
        blatt = self._mappe()
        erwartet = sum(len(s.elements) for s in self.registry.screens.values())
        self.assertEqual(blatt.max_row - 1, erwartet)

    def test_die_spalten_stimmen(self):
        blatt = self._mappe()
        kopf = [zelle.value for zelle in blatt[1]]
        self.assertEqual(kopf, [titel for _schluessel, titel in SPALTEN])

    def test_der_klartext_steht_daneben(self):
        """Sonst waere die Mappe nur eine Liste von Kuerzeln."""
        blatt = self._mappe()
        gefunden = False
        for zeile in blatt.iter_rows(min_row=2, values_only=True):
            if zeile[0] == "info_record_initial" and zeile[1] == "vendor":
                self.assertIn("Lieferantennummer", zeile[5])
                gefunden = True
        self.assertTrue(gefunden, "Feld info_record_initial.vendor fehlt")

    def test_der_pruefstand_wird_mitgeschrieben(self):
        self.registry.get("info_record_initial", "vendor").verified = True
        self.registry.get("info_record_initial", "material").verified = False
        blatt = self._mappe()
        stand = {(z[0], z[1]): z[7] for z in
                 blatt.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(stand[("info_record_initial", "vendor")], "ja")
        self.assertEqual(stand[("info_record_initial", "material")], "nein")

    def test_die_mappe_ist_zum_durchsehen_eingerichtet(self):
        """Eine Mappe, in der man sich verliert, wird nicht durchgesehen."""
        blatt = self._mappe()
        self.assertEqual(blatt.freeze_panes, "A2")
        self.assertIsNotNone(blatt.auto_filter.ref)

    def test_endung_wird_angelegt_wenn_der_ordner_fehlt(self):
        ziel = self.ordner / "tief" / "drin" / "ids.xlsx"
        self.assertTrue(exportiere_selektoren(self.registry, ziel).exists())


class RundlaufTest(unittest.TestCase):
    """Sichern, woanders einlesen -- der eigentliche Zweck."""

    def setUp(self):
        self.ordner = Path(tempfile.mkdtemp(prefix="rundlauf_"))

    def _aendern(self, pfad: Path, screen: str, feld: str, neue_id: str):
        from openpyxl import load_workbook

        mappe = load_workbook(pfad)
        blatt = mappe.active
        for zeile in blatt.iter_rows(min_row=2):
            if zeile[0].value == screen and zeile[1].value == feld:
                zeile[4].value = neue_id
        mappe.save(pfad)

    def test_unveraenderte_mappe_ergibt_keine_aenderung(self):
        quelle = SelectorRegistry()
        pfad = exportiere_selektoren(quelle, self.ordner / "ids.xlsx")
        ergebnis = lies_selektoren(SelectorRegistry(), pfad)
        self.assertFalse(ergebnis.hat_aenderungen)
        self.assertEqual(ergebnis.warnungen, [])

    def test_geaenderte_id_kommt_an(self):
        quelle = SelectorRegistry()
        pfad = exportiere_selektoren(quelle, self.ordner / "ids.xlsx")
        self._aendern(pfad, "info_record_initial", "vendor",
                      "wnd[0]/usr/subSUB0:SAPLMEGUI:0030/ctxtEINA-LIFNR")

        ziel = SelectorRegistry()
        ergebnis = lies_selektoren(ziel, pfad)
        self.assertIn(("info_record_initial", "vendor"), ergebnis.geaendert)
        alt, neu = ergebnis.geaendert[("info_record_initial", "vendor")]
        self.assertEqual(alt, "wnd[0]/usr/ctxtEINA-LIFNR")
        self.assertEqual(neu, "wnd[0]/usr/subSUB0:SAPLMEGUI:0030/ctxtEINA-LIFNR")

        uebernimm_selektoren(ziel, ergebnis)
        self.assertEqual(ziel.get("info_record_initial", "vendor").id, neu)

    def test_lesen_aendert_noch_nichts(self):
        """Erst anzeigen, dann uebernehmen -- nicht in einem Rutsch."""
        quelle = SelectorRegistry()
        pfad = exportiere_selektoren(quelle, self.ordner / "ids.xlsx")
        self._aendern(pfad, "info_record_initial", "vendor", "wnd[0]/usr/ctxtANDERS")

        ziel = SelectorRegistry()
        vorher = ziel.get("info_record_initial", "vendor").id
        lies_selektoren(ziel, pfad)
        self.assertEqual(ziel.get("info_record_initial", "vendor").id, vorher)


class FreigabeTest(unittest.TestCase):
    """Der Haken "Geprueft" -- das eigentlich Heikle."""

    def setUp(self):
        self.ordner = Path(tempfile.mkdtemp(prefix="freigabe_"))

    def test_haken_kommt_bei_unveraenderter_id_mit(self):
        quelle = SelectorRegistry()
        quelle.get("info_record_initial", "material").verified = True
        pfad = exportiere_selektoren(quelle, self.ordner / "ids.xlsx")

        ziel = SelectorRegistry()
        ziel.get("info_record_initial", "material").verified = False
        ergebnis = lies_selektoren(ziel, pfad)
        uebernimm_selektoren(ziel, ergebnis)
        self.assertTrue(ziel.get("info_record_initial", "material").verified)

    def test_geaenderte_id_ist_nie_schon_geprueft(self):
        """Der Import darf kein Weg sein, die Freigabe zu umgehen.

        Die Mappe kann aus einer anderen SAP-Anlage stammen.  Ob eine ID
        hier passt, weiss nur, wer sie an diesem System kontrolliert.
        """
        from openpyxl import load_workbook

        quelle = SelectorRegistry()
        quelle.get("info_record_initial", "vendor").verified = True
        pfad = exportiere_selektoren(quelle, self.ordner / "ids.xlsx")

        mappe = load_workbook(pfad)
        blatt = mappe.active
        for zeile in blatt.iter_rows(min_row=2):
            if (zeile[0].value == "info_record_initial"
                    and zeile[1].value == "vendor"):
                zeile[4].value = "wnd[0]/usr/ctxtFREMD-LIFNR"
                zeile[7].value = "ja"          # als geprueft ausgewiesen
        mappe.save(pfad)

        ziel = SelectorRegistry()
        ergebnis = lies_selektoren(ziel, pfad)
        uebernimm_selektoren(ziel, ergebnis)
        selektor = ziel.get("info_record_initial", "vendor")
        self.assertEqual(selektor.id, "wnd[0]/usr/ctxtFREMD-LIFNR")
        self.assertFalse(selektor.verified,
                         "Eine geaenderte ID darf nicht als geprueft gelten")

    def test_verschiedene_schreibweisen_fuer_ja(self):
        """Excel schreibt je nach Sprache und Format Verschiedenes."""
        from openpyxl import load_workbook

        for wert in ("ja", "Ja", "JA", "x", "X", "yes", "WAHR", "1", True):
            with self.subTest(wert=wert):
                quelle = SelectorRegistry()
                pfad = exportiere_selektoren(
                    quelle, self.ordner / f"ids_{str(wert).lower()}.xlsx")
                mappe = load_workbook(pfad)
                blatt = mappe.active
                for zeile in blatt.iter_rows(min_row=2):
                    if (zeile[0].value == "info_record_initial"
                            and zeile[1].value == "material"):
                        zeile[7].value = wert
                mappe.save(pfad)

                ziel = SelectorRegistry()
                ziel.get("info_record_initial", "material").verified = False
                ergebnis = lies_selektoren(ziel, pfad)
                uebernimm_selektoren(ziel, ergebnis)
                self.assertTrue(
                    ziel.get("info_record_initial", "material").verified)


class RobustheitTest(unittest.TestCase):
    """Was mit fremden oder beschaedigten Mappen passiert."""

    def setUp(self):
        self.ordner = Path(tempfile.mkdtemp(prefix="robust_"))
        self.registry = SelectorRegistry()

    def test_fremde_mappe_wird_erkannt(self):
        from openpyxl import Workbook

        mappe = Workbook()
        mappe.active.append(["Artikel", "Preis"])
        mappe.active.append(["4711", "2,95"])
        pfad = self.ordner / "angebot.xlsx"
        mappe.save(pfad)

        ergebnis = lies_selektoren(self.registry, pfad)
        self.assertFalse(ergebnis.hat_aenderungen)
        self.assertTrue(ergebnis.warnungen)

    def test_unbekannte_felder_werden_uebergangen(self):
        """Eine Mappe aus einer anderen Programmfassung."""
        from openpyxl import load_workbook

        pfad = exportiere_selektoren(self.registry, self.ordner / "ids.xlsx")
        mappe = load_workbook(pfad)
        mappe.active.append(["gibt_es_nicht", "auch_nicht", "", "", "wnd[0]/x",
                             "", "ja", "nein", ""])
        mappe.save(pfad)

        ergebnis = lies_selektoren(SelectorRegistry(), pfad)
        self.assertIn("gibt_es_nicht.auch_nicht", ergebnis.unbekannt)
        self.assertTrue(ergebnis.warnungen)
        self.assertFalse(ergebnis.geaendert)

    def test_umsortierte_spalten(self):
        """Gesucht wird ueber die Beschriftung, nicht ueber die Position."""
        from openpyxl import Workbook

        mappe = Workbook()
        blatt = mappe.active
        blatt.title = "SAP-Feld-IDs"
        blatt.append(["SAP-GUI-ID", "Feld", "Bildschirm", "Geprueft"])
        blatt.append(["wnd[0]/usr/ctxtNEU-LIFNR", "vendor",
                      "info_record_initial", "nein"])
        pfad = self.ordner / "umsortiert.xlsx"
        mappe.save(pfad)

        ergebnis = lies_selektoren(self.registry, pfad)
        self.assertEqual(
            ergebnis.geaendert.get(("info_record_initial", "vendor"))[1],
            "wnd[0]/usr/ctxtNEU-LIFNR")

    def test_leere_zeilen_stoeren_nicht(self):
        from openpyxl import load_workbook

        pfad = exportiere_selektoren(self.registry, self.ordner / "ids.xlsx")
        mappe = load_workbook(pfad)
        mappe.active.append([None] * len(SPALTEN))
        mappe.active.append(["", "", "", "", "", "", "", "", ""])
        mappe.save(pfad)
        ergebnis = lies_selektoren(SelectorRegistry(), pfad)
        self.assertFalse(ergebnis.hat_aenderungen)
        self.assertEqual(ergebnis.unbekannt, [])


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
