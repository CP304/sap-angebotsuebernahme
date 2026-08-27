"""SAP-Feld-IDs als Excel-Arbeitsmappe aus- und wieder einlesen.

Wozu
----
Die Feld-IDs sind muehsam erarbeitet: aufgezeichnet, zugeordnet, am
Zielsystem kontrolliert.  Sie liegen in einer JSON-Datei neben den
Einstellungen -- gut fuer das Programm, unhandlich fuer alles andere.

Als Arbeitsmappe lassen sie sich

* sichern, bevor jemand etwas umstellt,
* an Kollegen weitergeben, die dieselbe SAP-Anlage bedienen,
* auf einen zweiten Rechner uebertragen,
* und in Ruhe durchsehen -- mit Filtern, Kommentarspalte und dem
  Klartext zu jedem Feldnamen daneben.

Was hier NICHT passiert
-----------------------
Der Haken "Geprueft" wird beim Einlesen **nicht** blind uebernommen.
Geprueft heisst: jemand hat die ID am echten System kontrolliert.  Ob
das fuer diese Anlage gilt, weiss die Mappe nicht -- sie kann aus einem
anderen System stammen.  Uebernommen wird der Haken deshalb nur fuer
Felder, deren ID unveraendert bleibt; jede geaenderte ID gilt wieder als
ungeprueft und muss erneut bestaetigt werden.  Sonst waere der Import
ein Weg, die Freigabe zu umgehen.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

from ..sap.feldnamen import beschreibe_feld

logger = logging.getLogger(__name__)

__all__ = ["SPALTEN", "ExcelErgebnis", "exportiere_selektoren",
           "lies_selektoren", "uebernimm_selektoren"]

#: Die Spalten der Mappe.  Reihenfolge und Beschriftung sind Teil des
#: Formats -- beim Einlesen wird ueber die Beschriftung gesucht, nicht
#: ueber die Position, damit eine umsortierte Mappe weiter funktioniert.
SPALTEN = (
    ("bildschirm", "Bildschirm"),
    ("feld", "Feld"),
    ("transaktion", "Transaktion"),
    ("beschreibung", "Beschreibung"),
    ("id", "SAP-GUI-ID"),
    ("klartext", "So heisst das Feld in SAP"),
    ("pflicht", "Pflicht"),
    ("geprueft", "Geprueft"),
    ("bemerkung", "Bemerkung"),
)

#: Was in der Spalte "Geprueft" als Ja gilt -- Excel schreibt je nach
#: Sprache und Zellformat Verschiedenes.
_JA = {"ja", "j", "x", "yes", "y", "true", "wahr", "1"}


@dataclass
class ExcelErgebnis:
    """Was beim Einlesen herauskam -- zum Anzeigen, bevor etwas passiert."""

    geaendert: dict[tuple[str, str], tuple[str, str]] = field(default_factory=dict)
    """(Bildschirm, Feld) -> (bisherige ID, neue ID)"""
    freigaben: list[tuple[str, str]] = field(default_factory=list)
    """Felder, die durch die Mappe als geprueft gelten sollen."""
    unbekannt: list[str] = field(default_factory=list)
    """Zeilen, zu denen es hier kein Feld gibt (andere Programmfassung?)."""
    warnungen: list[str] = field(default_factory=list)

    @property
    def hat_aenderungen(self) -> bool:
        return bool(self.geaendert or self.freigaben)


def _openpyxl():
    """Erst beim Gebrauch laden -- ohne Excel laeuft alles Uebrige weiter."""
    try:
        import openpyxl
    except ImportError as fehler:  # pragma: no cover -- Paket ist Pflicht
        raise RuntimeError(
            "Fuer den Excel-Austausch fehlt das Paket 'openpyxl'. "
            "Abhilfe: pip install openpyxl") from fehler
    return openpyxl


def exportiere_selektoren(registry, ziel: Path) -> Path:
    """Alle Feld-IDs in eine Arbeitsmappe schreiben."""
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill

    mappe = openpyxl.Workbook()
    blatt = mappe.active
    blatt.title = "SAP-Feld-IDs"

    kopf = Font(bold=True, color="FFFFFF")
    fuellung = PatternFill("solid", fgColor="2F5597")
    for spalte, (_schluessel, titel) in enumerate(SPALTEN, start=1):
        zelle = blatt.cell(row=1, column=spalte, value=titel)
        zelle.font = kopf
        zelle.fill = fuellung
        zelle.alignment = Alignment(vertical="center")

    zeile = 2
    for screen_key, screen in registry.screens.items():
        for element_key, selector in screen.elements.items():
            werte = {
                "bildschirm": screen_key,
                "feld": element_key,
                "transaktion": screen.transaction,
                "beschreibung": selector.description,
                "id": selector.id,
                "klartext": beschreibe_feld(selector.id),
                "pflicht": "nein" if selector.optional else "ja",
                "geprueft": "ja" if selector.verified else "nein",
                "bemerkung": "",
            }
            for spalte, (schluessel, _titel) in enumerate(SPALTEN, start=1):
                blatt.cell(row=zeile, column=spalte, value=werte[schluessel])
            zeile += 1

    # Lesbar machen: Kopfzeile fixieren, filtern koennen, Spalten breit
    # genug.  Eine Mappe, in der man scrollen und suchen muss, wird nicht
    # durchgesehen -- und darum geht es hier.
    blatt.freeze_panes = "A2"
    blatt.auto_filter.ref = f"A1:{chr(64 + len(SPALTEN))}{zeile - 1}"
    breiten = {"bildschirm": 26, "feld": 26, "transaktion": 18,
               "beschreibung": 34, "id": 52, "klartext": 46,
               "pflicht": 9, "geprueft": 10, "bemerkung": 30}
    for spalte, (schluessel, _titel) in enumerate(SPALTEN, start=1):
        blatt.column_dimensions[chr(64 + spalte)].width = breiten[schluessel]

    ziel = Path(ziel)
    ziel.parent.mkdir(parents=True, exist_ok=True)
    mappe.save(ziel)
    logger.info("SAP-Feld-IDs nach Excel exportiert: %s (%d Zeilen)",
                ziel, zeile - 2)
    return ziel


def _kopfzeile(blatt) -> dict[str, int]:
    """Welche Spalte traegt welche Beschriftung?

    Gesucht wird ueber die Beschriftung, nicht ueber die Position: wer
    die Mappe umsortiert oder eine eigene Spalte einfuegt, soll sie
    trotzdem einlesen koennen.
    """
    nach_titel = {titel.strip().lower(): schluessel
                  for schluessel, titel in SPALTEN}
    gefunden: dict[str, int] = {}
    for spalte in range(1, blatt.max_column + 1):
        wert = blatt.cell(row=1, column=spalte).value
        if wert is None:
            continue
        schluessel = nach_titel.get(str(wert).strip().lower())
        if schluessel and schluessel not in gefunden:
            gefunden[schluessel] = spalte
    return gefunden


def lies_selektoren(registry, quelle: Path) -> ExcelErgebnis:
    """Eine Mappe lesen und mit dem aktuellen Stand vergleichen.

    Geaendert wird dabei noch nichts -- das Ergebnis ist zum Anzeigen
    gedacht, damit der Anwender sieht, was er uebernimmt.
    """
    openpyxl = _openpyxl()
    ergebnis = ExcelErgebnis()

    mappe = openpyxl.load_workbook(Path(quelle), data_only=True, read_only=True)
    try:
        blatt = (mappe["SAP-Feld-IDs"] if "SAP-Feld-IDs" in mappe.sheetnames
                 else mappe.active)
        spalten = _kopfzeile(blatt)
        fehlend = [titel for schluessel, titel in SPALTEN
                   if schluessel in ("bildschirm", "feld", "id")
                   and schluessel not in spalten]
        if fehlend:
            ergebnis.warnungen.append(
                "Die Mappe hat keine Spalte " + " und keine Spalte ".join(
                    f"„{t}“" for t in fehlend)
                + ". Stammt sie aus diesem Programm?")
            return ergebnis

        def lies(zeile, schluessel: str) -> str:
            spalte = spalten.get(schluessel)
            if not spalte:
                return ""
            wert = zeile[spalte - 1].value if spalte - 1 < len(zeile) else None
            return "" if wert is None else str(wert).strip()

        for zeile in blatt.iter_rows(min_row=2):
            screen_key = lies(zeile, "bildschirm")
            element_key = lies(zeile, "feld")
            if not screen_key or not element_key:
                continue
            screen = registry.screens.get(screen_key)
            if screen is None or element_key not in screen.elements:
                ergebnis.unbekannt.append(f"{screen_key}.{element_key}")
                continue

            selector = screen.elements[element_key]
            neue_id = lies(zeile, "id")
            geprueft = lies(zeile, "geprueft").lower() in _JA

            if neue_id and neue_id != selector.id:
                ergebnis.geaendert[(screen_key, element_key)] = (
                    selector.id, neue_id)
            elif geprueft and not selector.verified:
                # Nur wo die ID gleich bleibt, darf der Haken mitkommen.
                ergebnis.freigaben.append((screen_key, element_key))
    finally:
        mappe.close()

    if ergebnis.unbekannt:
        ergebnis.warnungen.append(
            f"{len(ergebnis.unbekannt)} Zeile(n) gehoeren zu Feldern, die es "
            "hier nicht gibt. Sie werden uebergangen -- stammt die Mappe aus "
            "einer anderen Programmfassung?")
    return ergebnis


def uebernimm_selektoren(registry, ergebnis: ExcelErgebnis) -> int:
    """Ein zuvor gelesenes Ergebnis tatsaechlich uebernehmen.

    Liefert die Anzahl geaenderter Felder.  Geaenderte IDs gelten danach
    als ungeprueft: ob eine ID aus einer fremden Mappe zu *dieser* Anlage
    passt, weiss nur, wer sie am System kontrolliert.
    """
    for (screen_key, element_key), (_alt, neu) in ergebnis.geaendert.items():
        registry.set_id(screen_key, element_key, neu)   # setzt verified=False
    for screen_key, element_key in ergebnis.freigaben:
        registry.get(screen_key, element_key).verified = True
    anzahl = len(ergebnis.geaendert) + len(ergebnis.freigaben)
    logger.info("Aus Excel uebernommen: %d Feld-ID(s) geaendert, %d freigegeben",
                len(ergebnis.geaendert), len(ergebnis.freigaben))
    return anzahl
